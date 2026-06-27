#!/usr/bin/env python3
"""Generate a B Hub payslip (xlsx + PDF) for one employee from a SALARY_<year>.xlsx month tab.

Fills only the month-variable cells of an employee payslip template, preserving the
employee's static master data (EPF no, tax no) and the template's layout/formulas.
Sources every figure from the salary sheet so manual copy-errors are corrected.

Usage:
  python generate_payslip.py --salary SALARY_2026.xlsx --month March \
      --employee "KEVIN EDDIE FRED" --template KEVIN_EDDIE_FRED.xlsx --out OUTDIR
"""
import argparse, os, re, sys, subprocess, json
import openpyxl

ONES = ["", "one", "two", "three", "four", "five", "six", "seven", "eight", "nine",
        "ten", "eleven", "twelve", "thirteen", "fourteen", "fifteen", "sixteen",
        "seventeen", "eighteen", "nineteen"]
TENS = ["", "", "twenty", "thirty", "forty", "fifty", "sixty", "seventy", "eighty", "ninety"]

def _three(n):
    w = []
    if n >= 100:
        w.append(ONES[n // 100] + " hundred"); n %= 100
    if n >= 20:
        t = TENS[n // 10]; t += "-" + ONES[n % 10] if n % 10 else ""; w.append(t)
    elif n > 0:
        w.append(ONES[n])
    return " ".join(w)

def int_to_words(n):
    if n == 0: return "zero"
    parts, scales = [], [(1_000_000, "million"), (1_000, "thousand"), (1, "")]
    for value, name in scales:
        if n >= value:
            chunk = n // value; n %= value
            parts.append(_three(chunk) + (" " + name if name else ""))
    return " ".join(parts).strip()

def money_in_words(amount):
    ringgit = int(round(amount * 100)) // 100
    sen = int(round(amount * 100)) % 100
    words = int_to_words(ringgit).title() + " Ringgit"
    if sen:
        words += " and " + int_to_words(sen).title() + " Sen"
    return words + " Only"

def norm_ic(v):
    return re.sub(r"\D", "", str(v)) if v is not None else ""

def norm_name(v):
    return re.sub(r"\s+", " ", str(v)).strip().lower() if v is not None else ""

def num(v):
    try: return float(v) if v not in (None, "") else 0.0
    except (TypeError, ValueError): return 0.0

def find_employee(ws, ident):
    hdr = {ws.cell(2, c).value: c for c in range(1, ws.max_column + 1)}
    cN, cIC = hdr.get("Name"), hdr.get("I/C")
    want_ic, want_nm = norm_ic(ident), norm_name(ident)
    by_ic = by_nm = None
    for r in range(3, ws.max_row + 1):
        nm = ws.cell(r, cN).value
        if nm is None or str(nm).strip() == "" or norm_name(nm) == "total :":
            continue
        if want_ic and norm_ic(ws.cell(r, cIC).value) == want_ic:
            by_ic = r
        if want_nm and norm_name(nm) == want_nm:
            by_nm = r
    return by_ic or by_nm, hdr

def generate(salary_path, month, employee, template_path, out_dir, year=None, epf_no=None, tax_no=None):
    flags = []
    if year is None:
        m = re.search(r"(20\d{2})", os.path.basename(salary_path))
        year = m.group(1) if m else ""
    wb_s = openpyxl.load_workbook(salary_path, data_only=False)
    if month not in wb_s.sheetnames:
        sys.exit(f"Month tab '{month}' not found. Tabs: {wb_s.sheetnames}")
    ws_s = wb_s[month]
    row, hdr = find_employee(ws_s, employee)
    if not row:
        sys.exit(f"Employee '{employee}' not found in {month}.")

    g = lambda h: ws_s.cell(row, hdr[h]).value
    name, ic, bank = g("Name"), g("I/C"), g("BANK ACCOUNT")
    D, E, F, G, H = num(g("SALARY")), num(g("Bonus/\nComm")), num(g("Unpaid Leave")), num(g("Deduction")), num(g("Allowance/Claim"))
    I, J, K, L = num(g("EPF (EE)")), num(g("SOCSO (EE)")), num(g("SOCSO EIS (EE)")), num(g("PCB"))
    O, P, Q = num(g("EPF (ER)")), num(g("SOCSO (ER)")), num(g("SOCSO EIS (ER)"))
    N = num(g("Advance"))
    nett = (D + H + E) - (F + G + I + J + K + L)            # mirrors salary col M
    m_col = num(g("TOTAL")) if not isinstance(g("TOTAL"), str) else None

    wb = openpyxl.load_workbook(template_path, data_only=False)
    ws = wb[wb.sheetnames[0]]
    MU = month.upper()

    ws["A5"] = f"PAYSLIP FOR {MU} {year}"
    ws["R7"] = f"END \u2013 {MU} {year}"
    ws["D7"], ws["D8"] = name, ic
    prev_bank = ws["R8"].value
    ws["R8"] = bank
    if prev_bank and norm_ic(prev_bank) != norm_ic(bank):
        flags.append(f"Bank corrected: template had '{prev_bank}', salary sheet says '{bank}'.")

    # earnings
    ws["A11"], ws["H11"] = "BASIC PAY", D
    er_row = 12
    for label, val in [("ALLOWANCE/CLAIM", H), ("BONUS/COMM", E)]:
        if val:
            ws.cell(er_row, 1).value = label; ws.cell(er_row, 8).value = val; er_row += 1
    # statutory deductions (labels already in template)
    ws["T11"], ws["T12"], ws["T13"], ws["T14"] = I, J, K, L
    de_row = 15
    for label, val in [("UNPAID LEAVE", F), ("OTHER DEDUCTION", G)]:
        if val:
            ws.cell(de_row, 11).value = label; ws.cell(de_row, 20).value = val; de_row += 1
    if N:
        flags.append(f"Advance RM{N:,.2f} present on salary sheet — not shown on payslip (matches net definition). Confirm if it should appear.")

    # bottom CURRENT MONTH employer figures (sourced from salary, fixes manual typos)
    ws["D26"], ws["F26"], ws["H26"] = O, P, Q
    # SOCSO member no derives from I/C digits; EPF/tax no come from the Employee Master
    ws["Q26"] = norm_ic(ic)
    if epf_no: ws["Q25"] = epf_no
    if tax_no: ws["Q27"] = tax_no
    if not ws["Q25"].value:
        flags.append("EPF member no missing — add it to the Employee Master for this employee.")
    if not ws["Q27"].value:
        flags.append("Income tax no missing — add it to the Employee Master for this employee.")
    # HR rule: PCB is the employee's monthly income tax, so show it in the bottom TAX (employee) line
    ws["J25"] = "=T14"

    # amount in words from net
    ws["E22"] = money_in_words(nett)

    os.makedirs(out_dir, exist_ok=True)
    safe = re.sub(r"[^A-Za-z0-9]+", "_", str(name)).strip("_")
    xlsx_out = os.path.join(out_dir, f"{safe}_{month}_{year}.xlsx")
    wb.save(xlsx_out)

    # recalc + verify
    recalc = os.path.join(os.path.dirname(os.path.abspath(__file__)), "recalc.py")
    if os.path.exists(recalc):
        subprocess.run([sys.executable, recalc, xlsx_out], capture_output=True, timeout=120)
    chk = openpyxl.load_workbook(xlsx_out, data_only=True)[wb.sheetnames[0]]
    nett_calc = num(chk["S21"].value)

    # PDF
    pdf_out = os.path.join(out_dir, f"{safe}_{month}_{year}.pdf")
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", out_dir, xlsx_out],
                   capture_output=True, timeout=180)
    produced_pdf = os.path.join(out_dir, f"{safe}_{month}_{year}.pdf")

    verify = {
        "employee": name, "month": month, "year": year,
        "net_python": round(nett, 2),
        "net_formula_in_file": round(nett_calc, 2),
        "net_matches_salary_TOTAL(M)": (m_col is None) or abs(nett - m_col) < 0.01,
        "salary_TOTAL(M)": m_col,
        "amount_in_words": ws["E22"].value,
        "earnings_total(D+H+E)": round(D + H + E, 2),
        "deductions_total(I+J+K+L+F+G)": round(I + J + K + L + F + G, 2),
        "pdf_created": os.path.exists(produced_pdf),
        "flags": flags,
    }
    return xlsx_out, produced_pdf, verify

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--salary", required=True)
    ap.add_argument("--month", required=True)
    ap.add_argument("--employee", required=True)
    ap.add_argument("--template", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--year", default=None)
    ap.add_argument("--epf-no", default=None)
    ap.add_argument("--tax-no", default=None)
    a = ap.parse_args()
    x, p, v = generate(a.salary, a.month, a.employee, a.template, a.out, a.year, a.epf_no, a.tax_no)
    print(json.dumps(v, indent=2))
    print("XLSX:", x)
    print("PDF :", p)
