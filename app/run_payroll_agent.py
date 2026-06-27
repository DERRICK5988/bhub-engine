#!/usr/bin/env python3
"""B Hub self-learning payslip agent.

Run it pointed at a folder (where Pensy drops SALARY_<year>.xlsx) — it:
  1. picks the latest SALARY_<year>.xlsx automatically (handles 2026 -> 2027 itself),
  2. picks the month to process (latest populated, or --month),
  3. detects full-timers = (has EPF/SOCSO/PCB this month) OR (already a confirmed full-timer
     in memory),  -- no manual indicator needed,
  4. LEARNS: records any newly-seen full-timer in the Employee Master (memory) so it is known
     next month even if a figure is blank,
  5. generates a payslip (xlsx + PDF) for every full-timer whose master data is complete,
  6. FLAGS new hires missing EPF/tax no, and full-timers expected but missing this month
     (possible resignation), without ever assuming,
  7. writes the updated memory back so the learning persists.

The memory file (employee_master.json) is human-auditable — open it to see what it learned.

Usage:
  python run_payroll_agent.py --drive-dir <folder> --template <blank payslip xlsx> \
      --memory employee_master.json --out <output dir> [--month June]
"""
import argparse, glob, json, os, re, sys
import openpyxl
from generate_payslip import generate, norm_ic, num

MONTHS = ["january", "february", "march", "april", "may", "june",
          "july", "august", "september", "october", "november", "december"]

def month_index(tab):
    t = re.sub(r"\(.*?\)", "", tab).strip().lower()[:3]          # ignore "(PT Route Setter)"
    aliases = {"feb": "february", "jan": "january"}              # tolerate the 'Feburary' typo
    for i, m in enumerate(MONTHS):
        if m.startswith(t) or aliases.get(t, "").startswith(m[:3]):
            return i
    return -1

def latest_salary_file(drive_dir):
    files = glob.glob(os.path.join(drive_dir, "SALARY_*.xlsx"))
    files = [f for f in files if re.search(r"SALARY_(20\d{2})", os.path.basename(f))
             and "Template" not in os.path.basename(f)]
    if not files:
        sys.exit(f"No SALARY_<year>.xlsx found in {drive_dir}")
    return max(files, key=lambda f: re.search(r"(20\d{2})", os.path.basename(f)).group(1))

def is_full_timer_row(ws, r, H):
    g = lambda h: ws.cell(r, H[h]).value if h in H else None
    return sum(num(g(h)) for h in ["EPF (EE)", "SOCSO (EE)", "SOCSO EIS (EE)", "PCB",
                                    "EPF (ER)", "SOCSO (ER)", "SOCSO EIS (ER)"]) > 0

def people_in_month(ws):
    H = {ws.cell(2, c).value: c for c in range(1, ws.max_column + 1)}
    out = []
    for r in range(3, ws.max_row + 1):
        nm = ws.cell(r, H["Name"]).value
        if not nm or str(nm).strip() == "" or str(nm).strip().lower().startswith("total"):
            continue
        out.append({"row": r, "name": str(nm).strip(),
                    "ic": str(ws.cell(r, H["I/C"]).value or "").strip(),
                    "ic_norm": norm_ic(ws.cell(r, H["I/C"]).value),
                    "bank": ws.cell(r, H["BANK ACCOUNT"]).value,
                    "salary": num(ws.cell(r, H["SALARY"]).value),
                    "statutory_ft": is_full_timer_row(ws, r, H)})
    return out, H

def pick_month(wb, requested):
    if requested:
        if requested not in wb.sheetnames:
            sys.exit(f"Month '{requested}' not in {wb.sheetnames}")
        return requested
    best, best_i = None, -1
    for tab in wb.sheetnames:
        if "(" in tab:                      # skip route-setter tabs for payslip month pick
            continue
        ws = wb[tab]; ppl, _ = people_in_month(ws)
        if any(p["statutory_ft"] for p in ppl) and month_index(tab) > best_i:
            best, best_i = tab, month_index(tab)
    if not best:
        sys.exit("Could not auto-pick a month with full-timer data; pass --month.")
    return best

def run(drive_dir, template, memory_path, out_dir, salary=None, month=None):
    salary = salary or latest_salary_file(drive_dir)
    year = re.search(r"(20\d{2})", os.path.basename(salary)).group(1)
    wb = openpyxl.load_workbook(salary, data_only=True)
    month = pick_month(wb, month)
    ppl, _ = people_in_month(wb[month])

    memory = json.load(open(memory_path)) if os.path.exists(memory_path) else {}
    learned, generated, flags = [], [], []

    # confirmed full-timers = statutory this month OR already known in memory
    full = [p for p in ppl if p["statutory_ft"]
            or (p["ic_norm"] in memory and str(memory[p["ic_norm"]].get("status", "active")).lower() != "resigned")]

    for p in full:
        ic = p["ic_norm"]
        if ic and ic not in memory:                              # LEARN a new full-timer
            memory[ic] = {"name": p["name"], "ic": p["ic"], "bank": p["bank"],
                          "status": "active", "epf_no": None, "tax_no": None,
                          "first_seen": f"{month} {year}"}
            learned.append(p["name"])
        m = memory.get(ic, {})
        if m.get("epf_no") and m.get("tax_no"):                  # complete -> generate
            x, pdf, v = generate(salary, month, p["name"], template, out_dir, year,
                                 epf_no=m["epf_no"], tax_no=m["tax_no"])
            folder, folder_id = m.get("payslip_folder"), m.get("payslip_folder_id")
            name_ok = (folder.strip().lower() == p["name"].strip().lower()) if folder else None
            generated.append({
                "name": p["name"], "pdf": os.path.basename(pdf), "net": v["net_python"],
                "target_folder": folder, "target_folder_id": folder_id,
                "name_matches_folder": name_ok,
                "safe_to_upload": bool(folder_id) and name_ok is not False,
                "flags": v["flags"]})
            if not folder_id:
                flags.append(f"FOLDER: {p['name']} has no locked Drive folder id. Create/confirm "
                             f"their Payslip subfolder, record its id as 'payslip_folder_id' in the "
                             f"Master, THEN upload. Never place this PDF in another person's folder.")
            elif name_ok is False:
                flags.append(f"STOP: folder name '{folder}' does not match employee '{p['name']}' "
                             f"— do NOT upload until the Master mapping is corrected.")
        else:                                                    # new hire / incomplete -> flag
            flags.append(f"ONBOARD: {p['name']} ({p['ic']}) is a full-timer but missing "
                         f"{'EPF no' if not m.get('epf_no') else ''} "
                         f"{'tax no' if not m.get('tax_no') else ''}".strip() +
                         " in the Employee Master — add it, then re-run.")

    # resignations / unexpected absences
    seen = {p["ic_norm"] for p in ppl if p["ic_norm"]}
    for ic, m in memory.items():
        if str(m.get("status", "active")).lower() == "active" and ic not in seen:
            flags.append(f"CHECK: {m.get('name')} is an active full-timer in memory but not in "
                         f"{month} {year} — resigned, on leave, or unpaid? Confirm.")

    json.dump(memory, open(memory_path, "w"), indent=2)          # persist what was learned

    return {"salary_file": os.path.basename(salary), "month": f"{month} {year}",
            "full_timers_detected": [p["name"] for p in full],
            "newly_learned_full_timers": learned,
            "payslips_generated": generated,
            "flags": flags,
            "memory_file": memory_path}

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--drive-dir", default=None)
    ap.add_argument("--salary", default=None)
    ap.add_argument("--template", required=True)
    ap.add_argument("--memory", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--month", default=None)
    a = ap.parse_args()
    if not a.drive_dir and not a.salary:
        sys.exit("Provide --drive-dir or --salary")
    print(json.dumps(run(a.drive_dir, a.template, a.memory, a.out, a.salary, a.month), indent=2))
