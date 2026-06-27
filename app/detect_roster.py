#!/usr/bin/env python3
"""Detect the payroll roster from a SALARY month tab and reconcile against the Employee Master.

Answers: who are the full-timers this month? who is NEW (needs onboarding)? who was expected
but is MISSING (possible resignation / unpaid month)? No names are hardcoded — it reads the data.

Usage:
  python detect_roster.py --salary SALARY_2026.xlsx --month March [--master employee_master.json]
"""
import argparse, json, re, sys
import openpyxl

def norm_ic(v): return re.sub(r"\D", "", str(v)) if v not in (None, "") else ""
def norm_name(v): return re.sub(r"\s+", " ", str(v)).strip().lower() if v not in (None, "") else ""
def num(v):
    try: return float(v) if v not in (None, "") else 0.0
    except (TypeError, ValueError): return 0.0

def detect(salary_path, month, master_path=None):
    wb = openpyxl.load_workbook(salary_path, data_only=True)
    if month not in wb.sheetnames:
        sys.exit(f"Month tab '{month}' not found. Tabs: {wb.sheetnames}")
    ws = wb[month]
    H = {ws.cell(2, c).value: c for c in range(1, ws.max_column + 1)}
    g = lambda r, h: ws.cell(r, H[h]).value if h in H else None

    full, part = [], []
    for r in range(3, ws.max_row + 1):
        nm = g(r, "Name")
        if not nm or str(nm).strip() == "" or str(nm).strip().lower().startswith("total"):
            continue
        statutory = sum(num(g(r, h)) for h in
                        ["EPF (EE)", "SOCSO (EE)", "SOCSO EIS (EE)", "PCB", "EPF (ER)", "SOCSO (ER)", "SOCSO EIS (ER)"])
        has_record = any(g(r, h) for h in ["START DATE", "CONFIRMATION DATE", "BEFORE PROBATION", "AFTER PROBATION"])
        person = {"row": r, "name": str(nm).strip(), "ic": str(g(r, "I/C") or "").strip(),
                  "ic_norm": norm_ic(g(r, "I/C")), "salary": num(g(r, "SALARY")),
                  "has_statutory": statutory > 0, "has_record": bool(has_record)}
        (full if (statutory > 0 or has_record) else part).append(person)

    report = {"month": month, "full_timers": full, "part_timers_count": len(part),
              "part_timer_names": [p["name"] for p in part]}

    if master_path:
        master = json.load(open(master_path))            # {ic_norm: {name,status,epf_no,tax_no,payslip_template}}
        det_ics = {p["ic_norm"] for p in full if p["ic_norm"]}
        new_hires = [p for p in full if p["ic_norm"] and p["ic_norm"] not in master]
        for p in new_hires:
            p["missing_master_data"] = ["epf_no", "tax_no", "payslip_template"]
        active_master = {ic: m for ic, m in master.items()
                         if str(m.get("status", "active")).lower() in ("active", "full-time", "full time")}
        missing = [{"ic": ic, **active_master[ic]} for ic in active_master if ic not in det_ics]
        resigned_but_present = [p for p in full if p["ic_norm"] in master
                                and str(master[p["ic_norm"]].get("status", "")).lower() == "resigned"]
        report.update({
            "known_full_timers": [p for p in full if p["ic_norm"] in master and p not in new_hires],
            "new_hires_need_onboarding": new_hires,
            "expected_but_missing_this_month": missing,   # possible resignation / unpaid / on leave
            "flagged_resigned_but_paid": resigned_but_present,
        })
    return report

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--salary", required=True)
    ap.add_argument("--month", required=True)
    ap.add_argument("--master", default=None)
    a = ap.parse_args()
    print(json.dumps(detect(a.salary, a.month, a.master), indent=2, default=str))
